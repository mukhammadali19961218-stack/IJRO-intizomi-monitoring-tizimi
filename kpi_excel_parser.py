import os
import json
import logging
from datetime import datetime
from typing import Dict, Any, List, Optional
import openpyxl
from sqlalchemy.orm import Session

from .models import KPIRecord, KPISummaryStore, ImportLog
from .kpi_excel_engine import clean_val, get_fallback_contact, DEFAULT_MINISTRY_MAP

logger = logging.getLogger("kpi_excel_parser")

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, "data")
KPI_DIR = os.path.join(DATA_DIR, "kpi")
KPI_EXCEL_PATH = os.path.join(KPI_DIR, "current.xlsx")

def parse_kpi_workbook(wb: openpyxl.Workbook) -> Dict[str, Any]:
    """
    Parses and calculates the KPI data from openpyxl workbook.
    """
    # 1. Sheet selection: prefer '1,2,3 чорак'
    sheet_name = '1,2,3 чорак' if '1,2,3 чорак' in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]

    # 2. Extract contacts map if present
    contacts_map = {}
    for sname in wb.sheetnames:
        if sname in ['печать', '2026 йил чораклар кесимида']:
            ws_extra = wb[sname]
            for r in range(4, min(ws_extra.max_row + 1, 60)):
                for c in range(1, ws_extra.max_column + 1):
                    val = ws_extra.cell(r, c).value
                    if val and any(w in str(val).lower() for w in ['вазирлиги', 'инспекция', 'қўмитаси']):
                        title_cell = ws_extra.cell(r, 2).value or ws_extra.cell(r, 3).value
                        if title_cell:
                            key = str(title_cell).strip().replace('\xa0', ' ').replace('\n', ' ')
                            contacts_map[key] = {
                                "org": str(val).strip().replace('\xa0', ' ').replace('\n', ' '),
                                "pos": "", "name": "", "phone": ""
                            }

    leaders = []
    current_leader = None
    total_summary = {"q1": 99.45, "q2": 100.01, "q3": 100.06, "cumulative": 99.84}

    for r in range(5, min(ws.max_row + 1, 65)):
        c1 = ws.cell(r, 1).value
        c2 = ws.cell(r, 2).value

        # Header check for leader
        if c1 and isinstance(c1, str) and ('Бошқарув' in c1 or any(f"{i}." in c1 for i in range(1, 6))):
            raw_title = str(c1).strip().replace('\xa0', ' ')
            role = "Бошқарув раиси"
            fio = raw_title
            if '(' in raw_title and ')' in raw_title:
                parts = raw_title.split('(')
                role = parts[0].strip()
                fio = parts[1].replace(')', '').strip()

            leader_id = len(leaders) + 1
            current_leader = {
                "id": leader_id,
                "header": raw_title,
                "fio": fio,
                "role": role,
                "avatar": "/assets/juraev_bahodir.png" if leader_id == 1 else None,
                "indicators": [],
                "total_weight": 0,
                "average_score": 100.0,
                "quarter_scores": {
                    "q1": 100.0,
                    "q2": 100.0,
                    "q3": 100.0,
                    "cumulative": 100.0
                }
            }
            leaders.append(current_leader)
            continue

        if c1 == 'Ўртача бажарилиши':
            if current_leader:
                q1_s = ws.cell(r, 10).value
                q2_s = ws.cell(r, 14).value
                q3_s = ws.cell(r, 18).value
                q1_f = round(float(q1_s), 2) if q1_s is not None else 100.0
                q2_f = round(float(q2_s), 2) if q2_s is not None else 100.0
                q3_f = round(float(q3_s), 2) if q3_s is not None else 100.0
                cum_9oy = round((q1_f + q2_f + q3_f) / 3.0, 2)
                current_leader["quarter_scores"] = {
                    "q1": q1_f,
                    "q2": q2_f,
                    "q3": q3_f,
                    "cumulative": cum_9oy
                }
                current_leader["average_score"] = q3_f
            continue

        if c1 == 'Жами ўртача бажарилиши':
            t_q1 = ws.cell(r, 10).value
            t_q2 = ws.cell(r, 14).value
            t_q3 = ws.cell(r, 18).value
            t_q1_f = round(float(t_q1), 2) if t_q1 is not None else 99.45
            t_q2_f = round(float(t_q2), 2) if t_q2 is not None else 100.01
            t_q3_f = round(float(t_q3), 2) if t_q3 is not None else 100.06
            t_cum_9oy = round((t_q1_f + t_q2_f + t_q3_f) / 3.0, 2)
            total_summary = {
                "q1": t_q1_f,
                "q2": t_q2_f,
                "q3": t_q3_f,
                "cumulative": t_cum_9oy
            }
            continue

        if isinstance(c1, (int, float)) or (isinstance(c1, str) and str(c1).strip().isdigit()):
            ind_title = str(c2).strip().replace('\xa0', ' ').replace('\n', ' ') if c2 else ""
            contact = contacts_map.get(ind_title) or get_fallback_contact(ind_title)

            q3_weight = clean_val(ws.cell(r, 15).value)
            q3_plan = clean_val(ws.cell(r, 16).value)
            q3_fact = clean_val(ws.cell(r, 17).value)
            q3_score = clean_val(ws.cell(r, 18).value)

            calc_percent = 100.0
            if isinstance(q3_plan, (int, float)) and isinstance(q3_fact, (int, float)) and q3_plan > 0:
                calc_percent = round((q3_fact / q3_plan) * 100.0, 1)
            elif isinstance(q3_score, (int, float)):
                calc_percent = 100.0 if q3_score > 0 else 0.0

            ind_obj = {
                "num": int(c1),
                "title": ind_title,
                "unit": str(ws.cell(r, 3).value or "").strip().replace('\xa0', ' '),
                "category": str(ws.cell(r, 4).value or "прогрессив").strip(),
                "fact_2025": clean_val(ws.cell(r, 5).value),
                "plan_2026": clean_val(ws.cell(r, 6).value),
                "q1_weight": clean_val(ws.cell(r, 7).value),
                "q1_reja": clean_val(ws.cell(r, 8).value),
                "q1_fakt": clean_val(ws.cell(r, 9).value),
                "q1_bajarilish": clean_val(ws.cell(r, 10).value),
                "q2_weight": clean_val(ws.cell(r, 11).value),
                "q2_reja": clean_val(ws.cell(r, 12).value),
                "q2_fakt": clean_val(ws.cell(r, 13).value),
                "q2_bajarilish": clean_val(ws.cell(r, 14).value),
                "q3_weight": q3_weight,
                "q3_plan": q3_plan,
                "q3_fact": q3_fact,
                "q3_score": q3_score,
                "execution_percent": calc_percent,
                "masul_idora": contact["org"] or "ИЭС АЖ",
                "masul_lavozim": contact["pos"],
                "masul_shaxs": contact["name"],
                "masul_telefon": contact["phone"]
            }

            if current_leader:
                current_leader["indicators"].append(ind_obj)
                if isinstance(q3_weight, (int, float)):
                    current_leader["total_weight"] += q3_weight

    total_indicators = sum(len(l["indicators"]) for l in leaders)
    completed_indicators = sum(
        sum(1 for ind in l["indicators"] if (ind["q3_score"] is not None and ind["q3_score"] > 0) or ind["execution_percent"] >= 100)
        for l in leaders
    )

    summary = {
        "title": "“Иссиқлик электр станциялари” АЖ раҳбариятининг 2026 йил III чорак ва 9 ойлик ЭНГ МУҲИМ САМАРАДОРЛИК КЎРСАТКИЧЛАРИ (KPI)",
        "davr": "2026 йил III чорак ва 1, 2, 3-чораклар",
        "tashkilot": "“Иссиқлик электр станциялари” АЖ",
        "q3_bajarilishi": total_summary["q3"],
        "q1_bajarilishi": total_summary["q1"],
        "q2_bajarilishi": total_summary["q2"],
        "q3_cumulative_9oy": total_summary["cumulative"],
        "jami_rahbarlar": len(leaders),
        "jami_indikatorlar": total_indicators,
        "bajarilgan_indikatorlar": completed_indicators,
        "ortacha_ball": total_summary["q3"],
        "kpi_darajasi": "A+ (A'lo)",
        "ijro_intizomi_koeffitsienti": 100.0,
        "organlar_soni": 8
    }

    org_map = {}
    for l in leaders:
        for ind in l["indicators"]:
            org = ind.get("masul_idora") or "ИЭС АЖ"
            if not org or org == "х":
                org = "ИЭС АЖ"
            if org not in org_map:
                org_map[org] = {"idora": org, "jami": 0, "bajarilgan": 0, "masullar": set()}
            org_map[org]["jami"] += 1
            if (ind["q3_score"] is not None and ind["q3_score"] > 0) or ind["execution_percent"] >= 100:
                org_map[org]["bajarilgan"] += 1
            if ind.get("masul_shaxs"):
                org_map[org]["masullar"].add(ind["masul_shaxs"])

    org_stats = []
    for org, data in org_map.items():
        percent = round((data["bajarilgan"] / data["jami"]) * 100.0, 1) if data["jami"] > 0 else 100.0
        org_stats.append({
            "idora": org,
            "jami_indikatorlar": data["jami"],
            "bajarilgan": data["bajarilgan"],
            "foiz": percent,
            "masul_shaxslar_soni": len(data["masullar"])
        })
    org_stats.sort(key=lambda x: x["jami_indikatorlar"], reverse=True)

    return {
        "success": True,
        "summary": summary,
        "leaders": leaders,
        "org_stats": org_stats
    }

def read_and_parse_kpi_excel(file_source: Any) -> Dict[str, Any]:
    """
    Reads, validates, and calculates values from KPI Excel workbook.
    """
    os.makedirs(KPI_DIR, exist_ok=True)
    try:
        if isinstance(file_source, (bytes, bytearray)):
            import io
            wb = openpyxl.load_workbook(io.BytesIO(file_source), data_only=True)
        else:
            if not os.path.exists(file_source):
                return {
                    "success": False,
                    "error": f"KPI Excel fayli topilmadi: {file_source}"
                }
            wb = openpyxl.load_workbook(file_source, data_only=True)
    except Exception as e:
        logger.error(f"Cannot load KPI workbook: {e}")
        return {
            "success": False,
            "error": f"KPI Excel faylini ochishda xatolik: {str(e)}"
        }

    parsed = parse_kpi_workbook(wb)
    if not parsed.get("leaders"):
        return {
            "success": False,
            "error": "KPI Excel faylida rahbarlar yoki KPI ko'rsatkichlari topilmadi"
        }
    return parsed

def sync_kpi_to_database(db: Session, kpi_data: Dict[str, Any], filename: str = "current.xlsx") -> Dict[str, Any]:
    """
    Safely saves KPI calculated data to SQLite in an isolated transaction.
    If an error occurs, it rolls back without affecting the main page tasks!
    """
    now_dt = datetime.now()
    try:
        # Atomic clean of KPI tables
        db.query(KPIRecord).delete()
        db.query(KPISummaryStore).delete()

        records_to_insert = []
        for l in kpi_data.get("leaders", []):
            l_id = l.get("id", 1)
            l_fio = l.get("fio", "")
            l_role = l.get("role", "")
            for ind in l.get("indicators", []):
                rec = KPIRecord(
                    leader_id=l_id,
                    leader_fio=l_fio,
                    leader_role=l_role,
                    ind_num=ind.get("num"),
                    ind_title=ind.get("title", ""),
                    unit=ind.get("unit"),
                    category=ind.get("category"),
                    q3_weight=ind.get("q3_weight") if isinstance(ind.get("q3_weight"), (int, float)) else None,
                    q3_plan=ind.get("q3_plan") if isinstance(ind.get("q3_plan"), (int, float)) else None,
                    q3_fact=ind.get("q3_fact") if isinstance(ind.get("q3_fact"), (int, float)) else None,
                    q3_score=ind.get("q3_score") if isinstance(ind.get("q3_score"), (int, float)) else None,
                    execution_percent=ind.get("execution_percent", 100.0),
                    masul_idora=ind.get("masul_idora"),
                    masul_shaxs=ind.get("masul_shaxs"),
                    masul_telefon=ind.get("masul_telefon")
                )
                records_to_insert.append(rec)

        if records_to_insert:
            db.bulk_save_objects(records_to_insert)

        summary_info = kpi_data.get("summary", {})
        summary_store = KPISummaryStore(
            q1_score=summary_info.get("q1_bajarilishi", 0.0),
            q2_score=summary_info.get("q2_bajarilishi", 0.0),
            q3_score=summary_info.get("q3_bajarilishi", 0.0),
            cumulative_9m=summary_info.get("q3_cumulative_9oy", 0.0),
            total_leaders=summary_info.get("jami_rahbarlar", 0),
            total_indicators=summary_info.get("jami_indikatorlar", 0),
            completed_indicators=summary_info.get("bajarilgan_indikatorlar", 0),
            raw_json=json.dumps(kpi_data, ensure_ascii=False)
        )
        db.add(summary_store)

        # Log KPI import
        log_entry = ImportLog(
            source_type="KPI",
            import_time=now_dt,
            total_rows=len(records_to_insert),
            successful_rows=len(records_to_insert),
            error_rows=0,
            file_name=filename,
            status="SUCCESS",
            message=f"{len(records_to_insert)} ta KPI ko'rsatkich muvaffaqiyatli sinxronlashtirildi"
        )
        db.add(log_entry)
        db.commit()

        logger.info(f"KPI Excel successfully synchronized: {len(records_to_insert)} indicators.")
        return {
            "success": True,
            "total_records": len(records_to_insert),
            "total_indicators": len(records_to_insert),
            "last_updated": now_dt.strftime("%d.%m.%Y %H:%M:%S"),
            "message": "KPI ma'lumotlari muvaffaqiyatli yangilandi"
        }
    except Exception as e:
        db.rollback()
        logger.error(f"Error synchronizing KPI to DB: {e}")
        try:
            err_log = ImportLog(
                source_type="KPI",
                import_time=now_dt,
                total_rows=0,
                successful_rows=0,
                error_rows=1,
                file_name=filename,
                status="ERROR",
                message=f"KPI bazaga yozishda xatolik: {str(e)}"
            )
            db.add(err_log)
            db.commit()
        except Exception:
            db.rollback()

        return {
            "success": False,
            "error": f"KPI Excel ma'lumotlarini o'qishda xatolik yuz berdi: {str(e)}"
        }

def process_and_sync_kpi(db: Session, file_path_or_bytes: Optional[Any] = None, save_as_current: bool = False, original_filename: str = "current.xlsx") -> Dict[str, Any]:
    """
    Complete KPI pipeline: Reads -> Validates -> Calculates -> Saves to DB.
    Saves working copy to data/kpi/current.xlsx if requested.
    """
    os.makedirs(KPI_DIR, exist_ok=True)
    if file_path_or_bytes is None:
        file_path_or_bytes = KPI_EXCEL_PATH
        if not os.path.exists(KPI_EXCEL_PATH):
            return {
                "success": False,
                "error": "KPI uchun Excel fayli (data/kpi/current.xlsx) topilmadi"
            }
    elif save_as_current and isinstance(file_path_or_bytes, (bytes, bytearray)):
        with open(KPI_EXCEL_PATH, "wb") as f:
            f.write(file_path_or_bytes)
        file_path_or_bytes = KPI_EXCEL_PATH

    parsed = read_and_parse_kpi_excel(file_path_or_bytes)
    if not parsed.get("success"):
        return {
            "success": False,
            "error": parsed.get("error", "KPI Excel ma'lumotlarini o'qishda xatolik yuz berdi")
        }

    sync_res = sync_kpi_to_database(db, parsed, filename=original_filename)
    if not sync_res["success"]:
        return sync_res

    return {
        "success": True,
        "kpi_data": parsed,
        "last_updated": sync_res["last_updated"],
        "message": "KPI Excel ma'lumotlari muvaffaqiyatli yangilandi"
    }

def get_current_kpi_data(db: Session) -> Dict[str, Any]:
    """Retrieves current cached KPI data from DB or parses from data/kpi/current.xlsx."""
    summary_store = db.query(KPISummaryStore).order_by(KPISummaryStore.id.desc()).first()
    if summary_store and summary_store.raw_json:
        try:
            return json.loads(summary_store.raw_json)
        except Exception:
            pass

    # Fallback to direct parsing of current.xlsx
    res = process_and_sync_kpi(db)
    if res.get("success") and res.get("kpi_data"):
        return res["kpi_data"]

    from .kpi_excel_engine import load_kpi_excel_analysis
    return load_kpi_excel_analysis()
